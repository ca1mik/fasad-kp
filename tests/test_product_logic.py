from copy import deepcopy
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from crm_backend.workspace import validate_finances
from tests.test_crm_portal import portal, PASSWORD, EDGE, sync_playwright
from tests.test_workspace_browser import sign_in


def test_financial_revision_is_atomic_and_balanced():
    base = {'clients': [{'id': 'c'}], 'orders': [{'id': 'o', 'clientId': 'c', 'amount': 1000, 'paidAmount': 0, 'payments': []}],
            'depositMovements': [{'id': 'top', 'clientId': 'c', 'kind': 'topup', 'amount': 500}]}
    validate_finances({}, base)
    changed = deepcopy(base)
    changed['orders'][0].update(paidAmount=400, payments=[{'id': 'p', 'method': 'deposit', 'amount': 400}])
    with pytest.raises(HTTPException):
        validate_finances(base, changed)
    changed['depositMovements'].append({'id': 's', 'clientId': 'c', 'orderId': 'o', 'kind': 'spend', 'amount': 400})
    validate_finances(base, changed)
    overdrawn = deepcopy(changed)
    overdrawn['depositMovements'].append({'id': 'w', 'clientId': 'c', 'kind': 'withdrawal', 'amount': 101, 'reason': 'Correction'})
    with pytest.raises(HTTPException):
        validate_finances(changed, overdrawn)
    duplicate = deepcopy(changed)
    duplicate['orders'][0]['payments'].append(duplicate['orders'][0]['payments'][0])
    with pytest.raises(HTTPException):
        validate_finances(changed, duplicate)
    removed = deepcopy(changed)
    removed['orders'][0].update(paidAmount=0, payments=[])
    removed['depositMovements'].append({'id': 'r', 'clientId': 'c', 'orderId': 'o', 'kind': 'refund', 'amount': 400})
    validate_finances(changed, removed)


def fill_lead(page, number='00072'):
    page.locator('#nav [data-route=workspace]').click()
    form = page.locator('#leadForm')
    for name, value in {'client': 'Тестовый заказчик', 'subject': 'Фасады', 'leadNumber': number}.items():
        form.locator(f'[name={name}]').fill(value)
    card = form.locator('.lead-part-card').first
    for name, value in {'partName': 'Фасад', 'partHeight': '700', 'partWidth': '150', 'partQty': '1'}.items():
        card.locator(f'[name={name}]').fill(value)
    return form, card


def test_server_draft_refresh_tabs_relogin_and_failed_submission(portal):
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=str(EDGE), headless=True)
        context = browser.new_context()
        page = context.new_page()
        page.set_default_timeout(7000)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.on('dialog', lambda dialog: dialog.accept())
        sign_in(page, portal)
        form, card = fill_lead(page)
        form.locator('[name=direction]').select_option('ПВХ')
        form.locator('[data-due-preset]').select_option('7w')
        due = form.locator('[name=due]').input_value()
        form.locator('.lead-optional').last.locator('summary').click()
        form.locator('[name=note]').fill('Не перепутать цвет\nСогласовать ручки')
        card.locator('.lead-part-options summary').click()
        card.locator('[name=partUnit]').select_option('м²')
        card.locator('[name=partPrice]').fill('1234.56')
        page.locator('#addLeadPart').click()
        form.locator('.lead-part-card').last.locator('[name=partName]').fill('Незаконченная строка')
        page.reload()
        page.locator("#workspaceExport").wait_for()
        form = page.locator('#leadForm')
        assert form.locator('[name=leadNumber]').input_value() == '00072'
        assert form.locator('[name=direction]').input_value() == 'ПВХ'
        assert form.locator('[name=due]').input_value() == due
        assert form.locator('[data-due-days]').input_value() == '7'
        assert form.locator('[name=partName]').last.input_value() == 'Незаконченная строка'
        assert form.locator('[name=partHeight]').last.input_value() == ''
        assert form.locator('[name=partPrice]').first.input_value() == '1234.56'
        assert form.locator('[name=partUnit]').first.input_value() == 'м²'
        # Duplicate the tab with inherited sessionStorage; each copy must fork.
        with page.expect_popup() as popup_info:
            page.evaluate('window.open(location.href)')
        other = popup_info.value
        other.locator('#workspaceExport').wait_for()
        other.locator('#leadForm [name=leadNumber]').fill('00073')
        other.locator('#leadForm [name=client]').fill('Другая вкладка')
        other.reload()
        other.locator("#workspaceExport").wait_for()
        assert other.locator('#leadForm [name=leadNumber]').input_value() == '00073'
        page.reload()
        page.locator("#workspaceExport").wait_for()
        assert page.locator('#leadForm [name=leadNumber]').input_value() == '00072'
        other.close()
        page.locator('#workspaceLogout').click()
        page.locator('#workspaceLogin [name=username]').fill('owner')
        page.locator('#workspaceLogin [name=password]').fill(PASSWORD)
        page.locator('#workspaceLogin button').click()
        page.locator('#workspaceExport').wait_for()
        assert form.locator('[name=leadNumber]').input_value() == '00072'
        # Finish the partial row, then abort the real server request.
        for name, value in {'partHeight': '500', 'partWidth': '300'}.items():
            form.locator(f'[name={name}]').last.fill(value)
        page.route('**/api/workspace', lambda route: route.abort() if route.request.method == 'PUT' else route.continue_())
        form.locator('[data-submit-mode=lead]').click()
        page.get_by_role('heading', name='Сохранение остановлено').wait_for()
        assert page.evaluate('Object.keys(localStorage).some(k=>k.includes("lead-draft:v1:")&&!k.endsWith(":tab"))')
        assert page.evaluate("fetch('/api/workspace').then(r=>r.json()).then(v=>v.revision)") == 0
        page.unroute('**/api/workspace')
        page.reload()
        page.locator("#workspaceExport").wait_for()
        assert form.locator('[name=leadNumber]').input_value() == '00072'
        form.locator('[data-submit-mode=lead]').click()
        page.wait_for_function('data.leads.length===1 && document.querySelector("#serverStatus").textContent.includes("версия 1")')
        assert page.evaluate('data.leads[0].id') == '00072-П'
        assert form.locator('[name=leadNumber]').input_value() == ''
        page.reload()
        page.locator("#workspaceExport").wait_for()
        assert form.locator('[name=leadNumber]').input_value() == ''
        assert not errors, errors
        browser.close()


def test_film_manual_layout_procurement_and_real_xlsx(portal, tmp_path):
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=str(EDGE), headless=True)
        page = browser.new_page(viewport={'width': 1600, 'height': 1100})
        page.set_default_timeout(7000)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        sign_in(page, portal)
        form, card = fill_lead(page)
        form.locator('[name=direction]').select_option('ПВХ')
        form.locator('[data-due-preset]').select_option('10w')
        due = form.locator('[name=due]').input_value()
        card.locator('[name=partHeight]').fill('2000')
        card.locator('[name=partWidth]').fill('1100')
        card.locator('[name=partColor]').fill('RAL 1019')
        card.locator('.lead-part-options summary').click()
        card.locator('[name=partPrice]').fill('100')
        page.locator('#addLeadPart').click()
        second = form.locator('.lead-part-card').last
        for name, value in {'partName': 'Большой фасад', 'partHeight': '2744', 'partWidth': '1100', 'partColor': 'RAL 9003'}.items():
            second.locator(f'[name={name}]').fill(value)
        note = 'Комментарий клиента: ручки согласовать.\n' + 'Длинное пожелание. ' * 35
        form.locator('.lead-optional').last.locator('summary').click()
        form.locator('[name=note]').fill(note)
        form.locator('[data-submit-mode=order]').click()
        page.locator('#newOrderForm').wait_for()
        assert page.locator('#newOrderForm [name=due]').input_value() == due
        page.locator('[data-modal-save]').click()
        page.wait_for_url('**/#orders/**')
        assert page.locator('.order-id h1').inner_text() == '00072-П'
        assert page.locator('#orderCustomerNote').inner_text() == note
        assert page.get_by_role('button', name='Пополнить депозит').count() == 0
        assert page.evaluate('FasadProduct.blankFile(data.orders[0]).length') > 1000
        with page.expect_download() as download_info:
            page.locator('[data-product=blank-xlsx]').click()
        output = tmp_path / 'blank.xlsx'
        download_info.value.save_as(output)
        ws = load_workbook(output).active
        comment_row = next(row[0].row for row in ws if row[0].value == 'Комментарий')
        assert ws.cell(comment_row, 2).value == note
        assert ws.cell(comment_row, 2).alignment.wrap_text is True
        assert ws.row_dimensions[comment_row].height > 32
        assert ws['F8'].number_format == '0.00'
        assert ws.cell(comment_row - 3, 1).value == 'Итого'
        page.locator('[data-action=open-cutting]').first.click()
        page.locator('#cuttingMaterialMode').select_option('film')
        page.locator('#filmOrderArea').wait_for()
        assert page.locator('#filmOrderArea').inner_text() == '7,09'
        assert page.evaluate('currentCuttingOrder().materialSpec.sheetWidth') == 3100
        page.locator('#filmColorFilter').select_option('RAL 1019')
        assert page.locator('#filmOrderArea').inner_text() == '7,09'
        assert page.locator('#cuttingMaps .cut-sheet-card:visible').count() == 1
        with page.expect_download() as supplier_info:
            page.locator('[data-product=film-xlsx]').click()
        supplier = load_workbook(BytesIO(supplier_info.value.path().read_bytes())).active
        assert supplier.cell(supplier.max_row, 2).value == pytest.approx(7.085)
        assert supplier.cell(supplier.max_row, 2).number_format == '0.00'
        page.locator('#openManualEditor').click()
        editor = page.locator('.layout-editor')
        editor.locator('.layout-piece[data-layout-key*="-P1#1"]').click()
        editor.locator('#editorX').fill('650')
        editor.locator('#editorY').fill('65')
        editor.locator('[data-edit=apply]').click()
        editor.locator('[data-edit=save]').click()
        page.locator('.layout-editor').wait_for(state='detached')
        assert page.locator('#filmOrderArea').inner_text() == '8,06'
        page.reload()
        page.locator("#workspaceExport").wait_for()
        assert page.locator('#filmOrderArea').inner_text() == '8,06'
        assert page.evaluate("data.orders[0].materialSpec.name") == 'МДФ 16 мм'
        page.locator('#cuttingMaterialMode').select_option('mdf')
        page.locator('#filmOrderArea').wait_for(state='detached')
        assert page.locator('#filmOrderArea').count() == 0
        assert not errors, errors
        browser.close()

def test_editing_source_lead_refreshes_draft_order_maps_quote_and_deadline(portal):
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path=str(EDGE),headless=True)
        page=browser.new_page();page.set_default_timeout(7000);sign_in(page,portal)
        form,card=fill_lead(page,'0089')
        card.locator('.lead-part-options summary').click();card.locator('[name=partPrice]').fill('100')
        form.locator('[data-submit-mode=order]').click();page.locator('#newOrderForm').wait_for()
        page.locator('[data-modal-save]').click();page.wait_for_url('**/#orders/0089')
        page.goto(portal+'/work/#leads/0089-%D0%AD');page.locator('#workspaceExport').wait_for()
        page.locator('[data-action=edit-lead]').click()
        edit=page.locator('#editLeadForm')
        for name,value in {'partHeight':'2700','partWidth':'1900','partQty':'3','partColor':'Новый цвет','due':'2027-01-20'}.items():edit.locator(f'[name={name}]').fill(value)
        page.locator('[data-modal-save]').click();page.locator('#editLeadForm').wait_for(state='hidden')
        state=page.evaluate('({order:data.orders.find(o=>o.id==="0089"),lead:data.leads[0]})')
        assert state['order']['parts'][0]['qty']==3
        assert state['order']['parts'][0]['color']=='Новый цвет'
        assert state['order']['amount']==300
        assert state['order']['dueDate']=='2027-01-20'
        assert state['lead']['parts'][0]['height']==2700
        page.goto(portal+'/work/#cutting/0089');page.locator('#workspaceExport').wait_for()
        assert page.locator('#cuttingMaps .cut-sheet-card').count()==3
        assert page.evaluate('calculateOrderNesting(currentCuttingOrder()).parts')==3
        browser.close()
