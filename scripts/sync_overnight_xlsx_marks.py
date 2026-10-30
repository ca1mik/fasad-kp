"""Mark only verified Gmail sends in the live CSV and overnight workbook."""

import csv
import io
import json
import os
from pathlib import Path
import sqlite3
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from serbia_outreach import worker_lock
import russia_outreach as ru

SOURCE = Path(r'F:\baza_proizvoditeley_fasadov_obzvon (1).xlsx')
OUT = ru.DATA / 'overnight-fasad-file-2026-10-01'


def main():
    with worker_lock(ru.DATA / 'worker.lock'):
        with sqlite3.connect(ru.DATA / 'campaign.sqlite3') as db:
            db.row_factory = sqlite3.Row
            deliveries = {r['email'].casefold(): dict(r)
                          for r in db.execute('SELECT lead_key,email,sent_at,status FROM deliveries')}
            receipts = {r['lead_key']: r['provider_id']
                        for r in db.execute('SELECT lead_key,provider_id FROM delivery_receipts')}
            suppressions = {r['email'].casefold(): r['reason']
                            for r in db.execute('SELECT email,reason FROM suppression')}
        queue = ru.SALES / 'campaign_leads.csv'
        with queue.open(encoding='utf-8-sig', newline='') as source:
            reader = csv.DictReader(source)
            fields = list(reader.fieldnames or [])
            rows = list(reader)
        if len(rows) != 96 or len({r['email'].casefold() for r in rows}) != 96:
            raise ValueError('Expected the reviewed 96-row queue')
        results = {}
        for row in rows:
            email = row['email'].casefold()
            delivery = deliveries.get(email)
            provider_id = receipts.get(delivery['lead_key']) if delivery else None
            verification = (ru.DATA / 'sent-verifications' / (provider_id + '.json')) if provider_id else None
            if delivery and delivery['status'] == 'committed' and verification.is_file():
                checked = json.loads(verification.read_text(encoding='utf-8'))
                if checked.get('provider_id') != provider_id:
                    raise ValueError('Gmail receipt and Sent verification differ')
                status, sent_at, gmail_id = 'sent', delivery['sent_at'], provider_id
            elif delivery:
                status, sent_at, gmail_id = 'unresolved', '', provider_id or ''
            elif email in suppressions:
                status, sent_at, gmail_id = 'excluded', '', ''
            else:
                status, sent_at, gmail_id = 'pending', '', ''
            row.update(send_status=status, sent_at_utc=sent_at, gmail_provider_id=gmail_id)
            results[email] = (status, sent_at, gmail_id)
        output = io.StringIO(newline='')
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
        new_csv = output.getvalue().encode('utf-8-sig')
        if queue.read_bytes() != new_csv:
            temp = queue.with_suffix('.csv.tmp')
            temp.write_bytes(new_csv)
            os.replace(temp, queue)
        workbook = load_workbook(SOURCE)
        sheet = workbook.worksheets[0]
        changed = False
        counts = {'sent': 0, 'pending': 0, 'excluded': 0, 'unresolved': 0,
                  'previously_sent': 0, 'no_email': 0}
        for i in range(5, sheet.max_row + 1):
            if not isinstance(sheet.cell(i, 1).value, int):
                continue
            email = str(sheet.cell(i, 10).value or '').strip().casefold()
            current = str(sheet.cell(i, 31).value or '')
            if current in {'Отправлено ранее', 'КП компании отправлено ранее'}:
                counts['previously_sent'] += 1
                continue
            if current == 'Нет email':
                counts['no_email'] += 1
                continue
            if email not in results:
                raise ValueError('Workbook email missing from campaign queue: ' + email)
            status, sent_at, provider_id = results[email]
            label = {'sent': 'Отправлено', 'pending': 'Ожидает отправки',
                     'excluded': 'Исключено', 'unresolved': 'Проверка отправки'}[status]
            desired = (label, sent_at, provider_id)
            current_values = tuple(sheet.cell(i, col).value or '' for col in (31, 32, 33))
            if current_values != desired:
                for col, value in zip((31, 32, 33), desired):
                    sheet.cell(i, col, value)
                changed = True
            counts[status] += 1
        if changed:
            temp = SOURCE.with_name(SOURCE.stem + '.sync.tmp.xlsx')
            workbook.save(temp)
            os.replace(temp, SOURCE)
        counts['campaign_sent_total'] = sum(r['send_status'] == 'sent' for r in rows)
        counts['source'] = str(SOURCE)
        (OUT / 'send-status.json').write_text(json.dumps(counts, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(counts, ensure_ascii=False))


if __name__ == '__main__':
    main()
