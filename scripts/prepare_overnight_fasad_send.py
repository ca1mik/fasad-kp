"""Add the owner's overnight facade producer file to the existing Russian queue."""

import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from serbia_outreach import worker_lock
import russia_outreach as ru

SOURCE = Path(r'F:\baza_proizvoditeley_fasadov_obzvon (1).xlsx')
SOURCE_SHA256 = 'fdf926f14dc57e221b6b9bea096b88685dc7cb83332a86301318ec76eb707937'
OUT = ru.DATA / 'overnight-fasad-file-2026-10-01'
SCOPE = 'Fasad Flow: ночной файл производителей фасадов от 2026-10-01, первое КП и демо'


def value(item):
    return str(item).strip() if item is not None else ''


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    with worker_lock(ru.DATA / 'worker.lock'):
        raw = SOURCE.read_bytes()
        if hashlib.sha256(raw).hexdigest() != SOURCE_SHA256:
            raise ValueError('Overnight workbook changed since review')
        workbook = load_workbook(SOURCE)
        sheet = workbook.worksheets[0]
        source_rows = [(i, [cell.value for cell in sheet[i]]) for i in range(5, sheet.max_row + 1)
                       if isinstance(sheet.cell(i, 1).value, int)]
        if len(source_rows) != 121 or [values[0] for _, values in source_rows] != list(range(1, 122)):
            raise ValueError('Expected the reviewed 121-row call list')
        queue = ru.SALES / 'campaign_leads.csv'
        old_bytes = queue.read_bytes()
        leads = ru.read_leads()
        if len(leads) != 54 or any(r.get('send_status') != 'sent' for r in leads):
            raise ValueError('Previous 54 Russian sends are not fully marked')
        old_emails = {r['email'].casefold() for r in leads}
        old_companies = {r['company'].strip().casefold() for r in leads}
        cfg_path = ru.SALES / 'campaign.json'
        config_bytes = cfg_path.read_bytes()
        cfg = ru.read_config()
        now = datetime.now(timezone.utc)
        candidates = []
        classifications = {}
        emails_in_file = set()
        for excel_row, cells in source_rows:
            email = value(cells[9]).casefold()
            company = value(cells[2])
            if not email:
                classifications[excel_row] = 'Нет email'
                continue
            if email in emails_in_file or not ru.EMAIL.fullmatch(email):
                raise ValueError('Duplicate or invalid email in owner file: ' + email)
            emails_in_file.add(email)
            if email in old_emails:
                classifications[excel_row] = 'Отправлено ранее'
            elif company.casefold() in old_companies:
                classifications[excel_row] = 'КП компании отправлено ранее'
            else:
                classifications[excel_row] = 'Ожидает отправки'
                candidates.append((excel_row, cells, email))
        if len(candidates) != 42 or len(emails_in_file) != 54:
            raise ValueError('Recipient counts changed since review')
        evidence = ru.SALES / 'owner-confirmation-overnight-2026-10-01.md'
        if evidence.exists():
            raise ValueError('Owner instruction is already recorded')
        lines = [
            '# Fasad Flow — поручение по ночному файлу производителей фасадов', '',
            f'Зарегистрировано: {now.isoformat()}. Это время регистрации поручения, не дата согласия получателей.',
            f'Исходный файл: {SOURCE}', f'SHA-256 исходного файла: {SOURCE_SHA256}', '',
            'После предыдущей рассылки владелец написал: «отправь все из файла вот из загрузок ночью которая».',
            'Файл найден по ярлыку недавно открытого документа от 01.10.2026 00:26; фактический путь сейчас на диске F:.',
            'Поручение относится к первому PDF-КП Fasad Flow по адресам ниже. Это разрешение владельца на отправку, '
            'а не независимо подтверждённое согласие получателей.', '',
            'В файле 121 компания: 67 без email, 11 адресов уже получили КП, ещё одна компания уже получила КП '
            'на другой адрес. Новых неповторяющихся адресатов 42. Отписки, стоп-листы, предыдущие обращения и '
            'лимиты продолжают действовать; обнаруженные при отправке повторы исключаются.', '',
            f'Область в конфигурации: {SCOPE}', '',
            '| № в файле | Компания | Email |', '|---|---|---|',
        ]
        lines += [f'| {cells[0]} | {value(cells[2])} | {email} |' for _, cells, email in candidates]
        evidence.write_text('\n'.join(lines) + '\n', encoding='utf-8')
        cfg['consent_scope'] = SCOPE
        header = list(leads[0])
        for excel_row, cells, email in candidates:
            company = value(cells[2])
            if '\r' in company or '\n' in company:
                raise ValueError('Invalid company header')
            product = value(cells[6])
            opening = ('В открытых источниках ваша компания указана среди производителей мебельных фасадов. '
                       'Обращаюсь по теме учёта заказов и этапов производства.')
            if product:
                opening = ('В открытых источниках указано ваше направление: ' + product.rstrip('. ') + '. '
                           'Предлагаю обсудить учёт заказов и этапов производства.')
            record = {key: '' for key in header}
            record.update(number=str(len(leads) + 1), company=company, city=value(cells[3]),
                          direction=value(cells[5]), email=email,
                          phone='; '.join(x for x in (value(cells[7]), value(cells[8])) if x),
                          website=value(cells[10]), source_url=value(cells[11]),
                          verified_at=now.isoformat(), email_verified='true',
                          fit_reason='Контакт из файла владельца; синтаксис email проверен, независимая проверка адреса не проводилась.',
                          origin='owner_overnight_fasad_file_2026_10_01', relationship='new_prospect_in_owner_file',
                          crm_status=value(cells[25]), opening_ru=opening,
                          benefit_ru=('Могу показать на тестовом заказе, как спецификация деталей связывается '
                                      'с расчётом, материалами, оплатой и этапами цеха.'),
                          approved='true', consent_at=now.isoformat(), consent_evidence=str(evidence),
                          consent_scope=SCOPE, do_not_contact='false', send_status='pending')
            if ru.consent_problem(record, cfg, now):
                raise ValueError('Owner instruction failed recipient validation: ' + email)
            leads.append(record)
        from copy import copy
        for col, title in [(31, 'КП: статус'), (32, 'КП: отправлено UTC'), (33, 'КП: Gmail ID')]:
            cell = sheet.cell(4, col, title)
            cell.font = copy(sheet.cell(4, 30).font)
            cell.fill = copy(sheet.cell(4, 30).fill)
            cell.alignment = copy(sheet.cell(4, 30).alignment)
        sheet.column_dimensions['AE'].width = 30
        sheet.column_dimensions['AF'].width = 24
        sheet.column_dimensions['AG'].width = 24
        for excel_row, _ in source_rows:
            sheet.cell(excel_row, 31, classifications[excel_row])
        backup = OUT / 'source-before-send.xlsx'
        backup.write_bytes(raw)
        (OUT / 'queue-before-send.csv').write_bytes(old_bytes)
        (OUT / 'config-before-send.json').write_bytes(config_bytes)
        output = io.StringIO(newline='')
        writer = csv.DictWriter(output, fieldnames=header)
        writer.writeheader()
        writer.writerows(leads)
        queue_temp = queue.with_suffix('.csv.tmp')
        config_temp = cfg_path.with_suffix('.json.tmp')
        workbook_temp = SOURCE.with_name(SOURCE.stem + '.tmp.xlsx')
        queue_temp.write_text(output.getvalue(), encoding='utf-8-sig', newline='')
        config_temp.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding='utf-8')
        workbook.save(workbook_temp)
        os.replace(workbook_temp, SOURCE)
        os.replace(config_temp, cfg_path)
        os.replace(queue_temp, queue)
        result = {'prepared_at': now.isoformat(), 'source': str(SOURCE), 'source_sha256': SOURCE_SHA256,
                  'rows': 121, 'new_queue': 42, 'previously_sent_email': 11,
                  'company_already_contacted_other_email': 1, 'missing_email': 67,
                  'queue_total': len(leads), 'evidence': str(evidence), 'sent_now': 0}
        (OUT / 'preparation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
